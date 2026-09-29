"""Official-source downloads and parsers. No generated or synthetic observations."""
from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from .provinces import BOT_TO_CODE, province_code_from_nesdc

BOT_URL = "https://app.bot.or.th/BTWS_STAT/statistics/BOTWEBSTAT.aspx?language=ENG&reportID=1008"
DOPA_URL = "https://stat.bora.dopa.go.th/new_stat/file/{yy}/stat_c{yy}.txt"
NESDC_URL = "https://www.nesdc.go.th/en/?p=107927&ddl=107925"
WB_URL = "https://api.worldbank.org/v2/country/THA/indicator/FP.CPI.TOTL.ZG"
POLICY_PAGE_URL = "https://www.bot.or.th/en/our-roles/monetary-policy/mpc-publication/policy-interest-rate.html"
RAW_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"
MONTHS = {m: i for i, m in enumerate("JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC".split(), 1)}


def _save_raw(source: str, data: bytes, extension: str) -> tuple[str, str]:
    digest = hashlib.sha256(data).hexdigest()
    folder = RAW_DIR / source
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{digest}.{extension}"
    if not path.exists():
        path.write_bytes(data)
    return str(path), digest


def _int(value: str) -> int | None:
    value = value.replace(",", "").strip()
    if value in ("", "-", "n.a.", "NA"):
        return None
    return int(round(float(value)))


def fetch_bot(start_year: int | None = None, start_month: int | None = None) -> tuple[list[dict], dict]:
    """BOT public table. POST request is its documented form, not an API claim."""
    session = requests.Session()
    response = session.get(BOT_URL, timeout=90)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    # The official page has hidden ASP.NET state; submit exactly its controls.
    newest = soup.select_one("#dgExcel tr")
    if newest is None:
        raise ValueError("BOT current-period table not found")
    current_label = newest.get_text(" ", strip=True).upper()
    current = re.search(r"\b([A-Z]{3})\s+(\d{4})\b", current_label)
    if current is None:
        raise ValueError("BOT current-period label not recognized")
    latest_year = int(current.group(2))
    latest_month = MONTHS[current.group(1)]
    if start_year is None or start_month is None:
        start_year, start_month = latest_year - 2, latest_month
    form = {i["name"]: i.get("value", "") for i in soup.select("input[name]")
            if i.get("type", "").lower() not in ("submit", "image")}
    form.update({"drpPeriod": "MTH", "drpFromMonth": f"xxxx{start_month:02d}xx",
                 "drpFromYear": f"{start_year}xxxx", "drpToMonth": f"xxxx{latest_month:02d}xx",
                 "drpToYear": f"{latest_year}xxxx", "btnSubmit": "Submit"})
    response = session.post(BOT_URL, data=form, timeout=180)
    response.raise_for_status()
    raw_path, digest = _save_raw("bot", response.content, "html")
    soup = BeautifulSoup(response.content, "html.parser")
    table = soup.select_one("#dgExcel")
    if table is None:
        raise ValueError("BOT table #dgExcel not found; source layout changed")
    rows = table.select("tr")
    headers = [c.get_text(" ", strip=True) for c in rows[0].select("td,th")]
    periods: list[date] = []
    for label in headers[2:]:
        match = re.search(r"\b([A-Z]{3})\s+(\d{4})\b", label.upper())
        if match:
            periods.append(date(int(match.group(2)), MONTHS[match.group(1)], 1))
    if len(periods) < 2:
        raise ValueError("BOT historical request returned fewer than two months; refusing to replace warehouse")
    observations: list[dict] = []
    seen = set()
    for row in rows[2:]:
        cells = [c.get_text(" ", strip=True) for c in row.select("td")]
        if len(cells) < 2:
            continue
        name = cells[1]
        if name not in BOT_TO_CODE:
            if name not in {"Head office", "Branches", "Central Region",
                            "Northeastern Region", "Northern Region", "Southern Region",
                            "Grand Total"}:
                raise ValueError(f"Unknown BOT row: {name}")
            continue
        code = BOT_TO_CODE[name]
        for idx, month in enumerate(periods):
            values = cells[2 + idx * 12:2 + (idx + 1) * 12]
            if len(values) != 12:
                raise ValueError(f"BOT missing columns for {name} {month}")
            key = (code, month)
            if key in seen:
                raise ValueError(f"BOT duplicate province-period: {key}")
            seen.add(key)
            observations.append({"province_code": code, "period": month,
                                 "branches": _int(values[0]),
                                 "deposits_million_baht": _int(values[6]),
                                 "credits_million_baht": _int(values[10])})
    if len(observations) != 77 * len(periods):
        raise ValueError(f"BOT expected 77 provinces per month, got {len(observations)}")
    return observations, {"url": BOT_URL, "raw_path": raw_path, "sha256": digest,
                          "periods": len(periods), "latest": str(max(periods))}


def fetch_dopa(years: tuple[int, ...] = (2023, 2024, 2025)) -> tuple[list[dict], list[dict]]:
    populations: list[dict] = []
    manifests: list[dict] = []
    for year in years:
        url = DOPA_URL.format(yy=str(year + 543)[2:])
        response = requests.get(url, timeout=90)
        response.raise_for_status()
        raw_path, digest = _save_raw("dopa", response.content, "txt")
        text = response.content.decode("utf-8-sig")
        year_rows = []
        for line in text.splitlines():
            values = line.split("|")
            if len(values) < 12 or not values[1].strip().isdigit():
                continue
            code = int(values[1].strip())
            if code == 0:
                continue
            if code not in BOT_TO_CODE.values():
                raise ValueError(f"Unknown DOPA province code: {code}")
            total = _int(values[11])
            if total is None or total <= 0:
                raise ValueError(f"Invalid DOPA population: {code} {year}")
            year_rows.append({"province_code": code, "reference_year": year,
                              "population": total,
                              "province_th": values[2].strip().removeprefix("จังหวัด")})
        if len(year_rows) != 77 or len({r["province_code"] for r in year_rows}) != 77:
            raise ValueError(f"DOPA expected 77 provinces for {year}, got {len(year_rows)}")
        populations.extend(year_rows)
        manifests.append({"url": url, "raw_path": raw_path, "sha256": digest,
                          "reference_year": year})
    return populations, manifests


def fetch_nesdc() -> tuple[list[dict], dict]:
    response = requests.get(NESDC_URL, timeout=120)
    response.raise_for_status()
    modified_header = response.headers.get("Last-Modified")
    if not modified_header:
        raise ValueError("NESDC download has no Last-Modified; publication date must be verified")
    source_last_modified = parsedate_to_datetime(modified_header).date()
    if not response.content.startswith(b"PK"):
        raise ValueError("NESDC response is not an XLSX workbook")
    raw_path, digest = _save_raw("nesdc", response.content, "xlsx")
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    sheet = workbook["PER CAPITA"]
    title = str(sheet.cell(2, 1).value)
    match = re.search(r"\b(20\d{2})", title)
    if not match:
        raise ValueError("NESDC reference year missing")
    year = int(match.group(1))
    output: list[dict] = []
    for row in sheet.iter_rows(values_only=True):
        if not (isinstance(row[0], (int, float)) and isinstance(row[1], str)
                and re.match(r"\d{4}\s", row[1])):
            continue
        code = province_code_from_nesdc(row[1])
        gpp = row[2]
        if not isinstance(gpp, (int, float)) or gpp <= 0:
            raise ValueError(f"Invalid NESDC GPP for {row[1]}")
        output.append({"province_code": code, "reference_year": year,
                       "gpp_million_baht": round(float(gpp), 3)})
    if len(output) != 77 or len({r["province_code"] for r in output}) != 77:
        raise ValueError(f"NESDC expected 77 provinces, got {len(output)}")
    return output, {"url": NESDC_URL, "raw_path": raw_path, "sha256": digest,
                    "reference_year": year, "last_modified": source_last_modified}


def fetch_world_bank() -> tuple[list[dict], dict]:
    """Real unauthenticated JSON API: Thailand annual CPI inflation (%)."""
    response = requests.get(WB_URL, params={"format": "json", "per_page": 100}, timeout=60)
    response.raise_for_status()
    raw_path, digest = _save_raw("worldbank", response.content, "json")
    metadata, observations = response.json()
    if metadata.get("pages") != 1 or len(observations) < 10:
        raise ValueError("Unexpected World Bank pagination or missing observations")
    last_updated = date.fromisoformat(metadata["lastupdated"])
    records = [{"reference_year": int(item["date"]), "inflation_pct": float(item["value"])}
               for item in observations if item.get("value") is not None]
    if len({r["reference_year"] for r in records}) != len(records):
        raise ValueError("World Bank duplicate reference year")
    return records, {"url": response.url, "raw_path": raw_path, "sha256": digest,
                     "last_updated": last_updated, "reference_year": max(r["reference_year"] for r in records)}


def fetch_policy_events() -> tuple[list[dict], dict]:
    """BOT's public MPC decision workbook; no API credentials required."""
    page = requests.get(POLICY_PAGE_URL, timeout=60)
    page.raise_for_status()
    soup = BeautifulSoup(page.content, "html.parser")
    links = [urljoin(POLICY_PAGE_URL, a["href"]) for a in soup.select("a[href]")
             if ".xlsx" in a["href"].lower() and "table-mpc" in a["href"].lower()]
    if len(links) != 1:
        raise ValueError(f"Expected one official policy workbook link, got {len(links)}")
    response = requests.get(links[0], timeout=90)
    response.raise_for_status()
    if not response.content.startswith(b"PK"):
        raise ValueError("BOT policy workbook is not XLSX")
    raw_path, digest = _save_raw("policy", response.content, "xlsx")
    sheet = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True).active
    output = []
    for row in sheet.iter_rows(values_only=True):
        if len(row) < 8 or not isinstance(row[2], datetime) or not isinstance(row[7], (int, float)):
            continue
        output.append({"event_date": row[2].date(), "policy_rate_pct": float(row[7]),
                       "decision": str(row[6] or "").strip()})
    if len(output) < 100 or len({r["event_date"] for r in output}) != len(output):
        raise ValueError("BOT policy workbook has too few events or duplicate dates")
    output.sort(key=lambda r: r["event_date"])
    return output, {"url": links[0], "raw_path": raw_path, "sha256": digest,
                    "latest": str(output[-1]["event_date"])}
